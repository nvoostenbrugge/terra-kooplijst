"""
Eenmalig overzetten van de oude kooplijst (Order Sheet .xlsx) naar deze app.

Alleen wat nog openstaat komt mee:
  - eerste kolom leeg        -> wacht op akkoord
  - eerste kolom "OK"        -> akkoord, nog bestellen
  - eerste kolom een datum   -> besteld, alleen als de datum in de laatste 14 dagen valt en
                                "Received" leeg is (alles wat ouder is blijft in het Excel-archief)
Regels met "geannuleerd" in de opmerking worden overgeslagen. Twee keer inladen maakt geen dubbelen.
"""
import re
from datetime import date, datetime, time, timedelta

from django.db import transaction
from django.utils import timezone

from . import logica
from .models import Bestelling, Regel

RECENT_DAGEN = 14
KOLOMMEN = {
    'omschrijving': ('description of goods', 'description'),
    'aantal': ('quantity',),
    'per': ('no.in package', 'no. in package', 'no.of packages'),
    'prijs': ('total cost', 'ttl. price'),
    'project': ('project name',),
    'opmerking': ('remarks',),
    'door': ('ordered by',),
    'link': ('link',),
    'ordernr': ('ordernr',),
    'ontvangen': ('received', 'receive date'),
    'factuur': ('invoice number',),
}


class ImportFout(Exception):
    pass


def _tekst(waarde):
    if waarde is None:
        return ''
    if isinstance(waarde, float) and waarde.is_integer():
        waarde = int(waarde)
    return str(waarde).strip()


def _datum(waarde, vandaag=None):
    """Besteldatum uit de eerste kolom.

    In de Google-sheet staan datums die als "06-10-2026" zijn getypt deels als echte datum opgeslagen,
    maar dan met dag en maand verwisseld (10 juni in plaats van 6 oktober). Voor zo'n cel nemen we
    daarom de verwisselde datum, tenzij die in de toekomst ligt. Tekst als "29-09-2026" is gewoon dd-mm-jjjj.
    """
    if isinstance(waarde, datetime):
        waarde = waarde.date()
    if isinstance(waarde, date):
        vandaag = vandaag or timezone.localdate()
        try:
            gewisseld = date(waarde.year, waarde.day, waarde.month)
        except ValueError:
            return waarde
        return gewisseld if gewisseld <= vandaag else waarde
    m = re.fullmatch(r'(\d{1,2})-(\d{1,2})-(\d{4})', _tekst(waarde))
    if not m:
        return None
    try:
        return date(int(m[3]), int(m[2]), int(m[1]))
    except ValueError:
        return None


def _kies_blad(wb):
    jaar = str(timezone.localdate().year)
    for ws in wb.worksheets:
        if jaar in ws.title:
            return ws
    return wb.worksheets[0]


def _kopregel(ws):
    for nr, rij in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
        namen = [_tekst(c).lower() for c in rij]
        if any(n in KOLOMMEN['omschrijving'] for n in namen):
            kolom = {}
            for sleutel, opties in KOLOMMEN.items():
                kolom[sleutel] = next((i for i, n in enumerate(namen) if n in opties), None)
            return nr, kolom
    raise ImportFout('Dit lijkt niet de kooplijst: de kolom "Description of Goods" ontbreekt.')


def importeer_xlsx(bestand, proef=False):
    """Leest het Excel-bestand en maakt regels aan. Met proef=True wordt er niets opgeslagen."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(bestand, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 - elk onleesbaar bestand geeft dezelfde melding
        raise ImportFout('Het bestand kon niet gelezen worden. Kies het .xlsx-bestand van de kooplijst.') from e
    ws = _kies_blad(wb)
    kop, kolom = _kopregel(ws)
    vandaag = timezone.localdate()
    bestaand = set(Regel.objects.filter(geimporteerd=True).values_list('omschrijving', 'link', 'aantal'))
    tel = {'aangevraagd': 0, 'goedgekeurd': 0, 'besteld': 0, 'overgeslagen': 0, 'dubbel': 0}
    nieuw, bestellingen = [], {}

    def cel(rij, sleutel):
        i = kolom.get(sleutel)
        return rij[i] if i is not None and i < len(rij) else None

    for rij in ws.iter_rows(min_row=kop + 1, values_only=True):
        if not rij:
            continue
        omschrijving, link = _tekst(cel(rij, 'omschrijving'))[:500], _tekst(cel(rij, 'link'))[:2000]
        if not omschrijving and not logica.nette_link(link):
            continue
        eerste, opmerking = rij[0], _tekst(cel(rij, 'opmerking'))[:500]
        besteld_op = _datum(eerste)
        if re.search(r'geannuleerd|cancel', opmerking, re.I):
            tel['overgeslagen'] += 1
            continue
        if _tekst(eerste) == '':
            status = Regel.AANGEVRAAGD
        elif _tekst(eerste).lower() == 'ok':
            status = Regel.GOEDGEKEURD
        elif besteld_op and (vandaag - besteld_op).days <= RECENT_DAGEN and besteld_op <= vandaag and not _tekst(cel(rij, 'ontvangen')):
            status = Regel.BESTELD
        else:
            tel['overgeslagen'] += 1
            continue
        try:
            aantal = max(1, int(float(re.match(r'[\d.]+', _tekst(cel(rij, 'aantal')) or '1')[0])))
        except (TypeError, ValueError):
            aantal = 1
        try:
            per = max(1, int(float(re.match(r'[\d.]+', _tekst(cel(rij, 'per')) or '1')[0])))
        except (TypeError, ValueError):
            per = 1
        if (omschrijving, link, aantal) in bestaand:
            tel['dubbel'] += 1
            continue
        bestaand.add((omschrijving, link, aantal))
        regel = Regel(
            geimporteerd=True, omschrijving=omschrijving, link=link, winkel=logica.winkel_van(link), aantal=aantal, per_verpakking=per,
            prijs=logica.naar_prijs(cel(rij, 'prijs')), project=_tekst(cel(rij, 'project'))[:200], opmerking=opmerking,
            aanvrager_naam=_tekst(cel(rij, 'door'))[:200].capitalize() if _tekst(cel(rij, 'door')).islower() else _tekst(cel(rij, 'door'))[:200],
            spoed=bool(re.search(r'spoed', opmerking, re.I)), status=status, factuurnr=_tekst(cel(rij, 'factuur'))[:200],
        )
        if status == Regel.GOEDGEKEURD:
            regel.goedgekeurd_op = timezone.now()
        if status == Regel.BESTELD:
            regel._order = (regel.winkel, _tekst(cel(rij, 'ordernr'))[:200], besteld_op)
        tel[status] += 1
        nieuw.append(regel)

    if not proef:
        with transaction.atomic():
            for regel in nieuw:
                sleutel = getattr(regel, '_order', None)
                if sleutel:
                    if sleutel not in bestellingen:
                        moment = timezone.make_aware(datetime.combine(sleutel[2], time(14, 0)))
                        bestellingen[sleutel] = Bestelling.objects.create(winkel=sleutel[0], ordernr=sleutel[1], besteld_op=moment)
                    regel.bestelling = bestellingen[sleutel]
                regel.save()
    return {'blad': ws.title, 'proef': proef, 'geteld': tel, 'totaal': len(nieuw)}
